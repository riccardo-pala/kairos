"""Build the stakeholder DOCX from the single functional Markdown source.

Requires python-docx. Run with the Codex bundled Python or a separate docs venv.
Product dependencies are intentionally kept independent of document tooling.
"""

import argparse
import re
from datetime import date
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]


def inline(paragraph, text):
    text = re.sub(r"\[([^]]+)\]\([^)]+\)", r"\1", text)
    for piece in re.split(r"(\*\*.*?\*\*)", text):
        run = paragraph.add_run(piece.strip("*") if piece.startswith("**") else piece)
        run.bold = piece.startswith("**")


def table(doc, rows):
    result = doc.add_table(rows=0, cols=len(rows[0]))
    result.alignment = WD_TABLE_ALIGNMENT.CENTER
    result.autofit = False
    widths = [1.5, 5.2]
    for column, width in zip(result.columns, widths):
        column.width = Inches(width)
    for index, values in enumerate(rows):
        row = result.add_row()
        no_split = OxmlElement("w:cantSplit")
        row._tr.get_or_add_trPr().append(no_split)
        if index == 0:
            repeat = OxmlElement("w:tblHeader")
            row._tr.get_or_add_trPr().append(repeat)
        for col, text in enumerate(values):
            cell = row.cells[col]
            cell.width = Inches(widths[col])
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            props = cell._tc.get_or_add_tcPr()
            borders = OxmlElement("w:tcBorders")
            for edge in ("top", "left", "bottom", "right"):
                item = OxmlElement(f"w:{edge}")
                for key, value in {"val": "single", "sz": "4", "color": "D9D9D9"}.items():
                    item.set(qn(f"w:{key}"), value)
                borders.append(item)
            props.append(borders)
            margins = OxmlElement("w:tcMar")
            for edge in ("top", "left", "bottom", "right"):
                item = OxmlElement(f"w:{edge}")
                item.set(qn("w:w"), "110")
                item.set(qn("w:type"), "dxa")
                margins.append(item)
            props.append(margins)
            shade = OxmlElement("w:shd")
            shade.set(qn("w:fill"), "E9EFF5" if index == 0 else "FFFFFF")
            props.append(shade)
            paragraph = cell.paragraphs[0]
            paragraph.paragraph_format.space_after = Pt(2)
            inline(paragraph, text)
            for run in paragraph.runs:
                run.font.size = Pt(10.5)
                if index == 0:
                    run.bold = True
    doc.add_paragraph().paragraph_format.space_after = Pt(3)


def render_table(doc, rows):
    # Narrative matrices are easier to read as paragraphs in a portrait report.
    if rows[0][0] == "ID":
        for identifier, requirement, status in rows[1:]:
            paragraph = doc.add_paragraph()
            paragraph.paragraph_format.space_before = Pt(8)
            paragraph.paragraph_format.space_after = Pt(3)
            paragraph.paragraph_format.keep_with_next = True
            inline(paragraph, f"{identifier}  {requirement}")
            paragraph = doc.add_paragraph()
            inline(paragraph, f"**Status**  {status}")
            for run in paragraph.runs:
                run.font.size = Pt(10)
    elif rows[0][0] == "Area":
        for area, target, current, future in rows[1:]:
            doc.add_heading(area.replace(",", ""), level=2)
            for label, content in (("Goal", target), ("Current availability", current), ("Future work", future)):
                inline(doc.add_paragraph(), f"**{label}**  {content}")
    else:
        table(doc, rows)


def build(output, revision, updated):
    source = ROOT / "docs/functional/requirements.md"
    doc = Document()
    section = doc.sections[0]
    section.page_width, section.page_height = Inches(8.5), Inches(11)
    section.top_margin = section.bottom_margin = Inches(0.8)
    section.left_margin = section.right_margin = Inches(0.9)
    section.different_first_page_header_footer = True
    styles = doc.styles
    for name in ("Normal", "Title", "Subtitle", "Heading 1", "Heading 2", "Heading 3"):
        styles[name].font.name = "Calibri"
        styles[name].font.color.rgb = RGBColor(0, 0, 0)
        for border in styles[name].element.findall(".//" + qn("w:pBdr")):
            border.getparent().remove(border)
    normal = styles["Normal"]
    normal.font.size = Pt(11)
    normal.paragraph_format.line_spacing = 1.12
    normal.paragraph_format.space_after = Pt(7)
    for name, size in (("Heading 1", 19), ("Heading 2", 13)):
        styles[name].font.size = Pt(size)
        styles[name].paragraph_format.space_before = Pt(15)
        styles[name].paragraph_format.space_after = Pt(7)
    styles["Title"].font.size = Pt(36)

    cover = doc.add_paragraph()
    cover.paragraph_format.space_before = Pt(35)
    cover.alignment = WD_ALIGN_PARAGRAPH.CENTER
    picture = cover.add_run().add_picture(str(ROOT / "assets/logo.png"), width=Inches(2.5))
    picture._inline.docPr.set("descr", "Kairos logo")
    p = doc.add_paragraph("Kairos", style="Title")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p = doc.add_paragraph("Product Vision and Functional Requirements", style="Subtitle")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(30)
    p = doc.add_paragraph("A personal assistant for people who procrastinate")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p = doc.add_paragraph("Product overview and first proof-of-concept scope")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(35)
    months = ("January", "February", "March", "April", "May", "June", "July", "August", "September", "October", "November", "December")
    date_label = f"{updated.day} {months[updated.month - 1]} {updated.year}"
    for text in (f"Revision {revision}", date_label, "Android  |  English and Italian"):
        p = doc.add_paragraph(text)
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.add_page_break()

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    footer.add_run(f"Kairos  ·  {revision}  ·  ")
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    footer._p.append(field)
    for run in footer.runs:
        run.font.size = Pt(9)

    lines = source.read_text().splitlines()
    section_number = 0
    i = 0
    while i < len(lines):
        line = lines[i].strip()
        i += 1
        if not line or line.startswith("# ") or line.startswith("**Document status:"):
            continue
        if line.startswith("## "):
            heading = line[3:]
            if heading == "Document management":
                break  # Internal publishing policy is outside the stakeholder report.
            section_number += 1
            heading = {"Executive summary": "Executive summary", "Proof-of-concept scope and future work": "Proof-of-concept scope and future work"}.get(heading, heading)
            paragraph = doc.add_heading(f"{section_number}  {heading}", level=1)
            if heading in ("Functional requirements", "Open questions"):
                paragraph.paragraph_format.page_break_before = True
        elif line.startswith("|"):
            rows = []
            while True:
                values = [value.strip() for value in line.strip("|").split("|")]
                if not all(re.fullmatch(r":?-+:?", value) for value in values):
                    rows.append(values)
                if i >= len(lines) or not lines[i].strip().startswith("|"):
                    break
                line = lines[i].strip()
                i += 1
            render_table(doc, rows)
        else:
            if line.startswith("- "):
                paragraph = doc.add_paragraph(style="List Bullet")
                line = line[2:]
            elif re.match(r"^\d+\. ", line):
                paragraph = doc.add_paragraph()
                paragraph.paragraph_format.left_indent = Inches(0.24)
                paragraph.paragraph_format.first_line_indent = Inches(-0.24)
                paragraph.paragraph_format.space_after = Pt(2)
            else:
                paragraph = doc.add_paragraph()
            inline(paragraph, line)

    doc.core_properties.title = "Kairos Product Vision and Functional Requirements"
    doc.core_properties.subject = "Product overview and proof-of-concept scope"
    doc.core_properties.author = "Kairos"
    doc.core_properties.version = revision
    output.parent.mkdir(parents=True, exist_ok=True)
    doc.save(output)
    print(output)


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--output", type=Path, default=ROOT / "docs/functional/Kairos_Functional_Requirements.docx")
    parser.add_argument("--revision", default="0.3")
    parser.add_argument("--date", type=date.fromisoformat, default=date(2026, 10, 25))
    args = parser.parse_args()
    build(args.output, args.revision, args.date)
